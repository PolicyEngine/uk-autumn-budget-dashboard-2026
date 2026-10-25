"""Read one forecast's quarterly real household disposable income per person.

Each Economic and fiscal outlook's economy tables (called "supplementary"
tables before 2024, "detailed forecast" tables since) have a Table 1.5,
"Per capita" or "Per person", with two blocks of columns: "By total
population" and "Ages 16+". In the first block, the "Real household
disposable income" column under the heading "Index: 2008Q1=100" is RHDI
(the ONS's chained volume measure, NRJR) divided by total population (ONS
EBAQ), indexed to 2008 Q1. The quarterly rows (``2008Q1`` onwards) run
from 2008 to the forecast's last quarter.

The sheet, block, column and rows are found by their headings and labels,
never by fixed positions; a sheet without them raises :class:`EfoTableError`.
The table does not say which quarters are outturn: the build decides each
forecast's launch quarter from its publication date
(:func:`uk_budget_data.forecast_vintages.periods.quarterly_launch`).
"""

import io
import re
import zipfile
from pathlib import Path

import openpyxl

QUARTER = re.compile(r"^(\d{4})\s*Q([1-4])$")


class EfoTableError(ValueError):
    """An EFO workbook without the per-person table the reader expects."""


def _text(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _open(source: Path, member: str = None):
    data = Path(source).read_bytes()
    if member:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            data = archive.read(member)
    return openpyxl.load_workbook(
        io.BytesIO(data), read_only=True, data_only=True
    )


def _find_sheet(workbook):
    for ws in workbook.worksheets:
        if ws.title.strip() != "1.5":
            continue
        head = [
            _text(c)
            for r in ws.iter_rows(max_row=3, values_only=True)
            for c in r
        ]
        if any(re.search(r"per (capita|person)", h, re.I) for h in head):
            return ws
    raise EfoTableError("No Table 1.5 (per capita / per person)")


def read_rhdi_per_person(source: Path, member: str = None) -> dict:
    """Quarterly RHDI per person index (2008Q1=100), by total population.

    ``member`` names the workbook inside ``source`` when ``source`` is a
    zip. Returns ``{"2008Q1": 100.0, ...}``.
    """
    ws = _find_sheet(_open(source, member))
    rows = list(ws.iter_rows(max_row=200, values_only=True))
    header = rows[:8]

    def find(pattern, row_from=0):
        for i, row in enumerate(header[row_from:], start=row_from):
            for j, cell in enumerate(row):
                if re.search(pattern, _text(cell), re.I):
                    return i, j
        return None

    block = find(r"^by total population$")
    other = find(r"^ages 16\+$")
    if block is None or other is None:
        raise EfoTableError("Table 1.5 has no 'By total population' block")
    first, last = block[1], other[1]
    if last <= first:
        raise EfoTableError("Table 1.5's blocks are not in the expected order")

    # Within the block: the index heading row, then the series names row
    index_head = find(r"^index: 2008\s*q1\s*=\s*100$")
    if index_head is None or not first <= index_head[1] < last:
        raise EfoTableError("Table 1.5 has no 2008Q1=100 index in its block")
    names_row = None
    column = None
    for i in range(index_head[0] + 1, len(header)):
        for j in range(first, last):
            cell = header[i][j] if j < len(header[i]) else None
            # October 2024 marks the heading with a footnote number
            if re.fullmatch(
                r"real household disposable income\d*", _text(cell), re.I
            ):
                names_row, column = i, j
                break
        if column is not None:
            break
    if column is None:
        raise EfoTableError("Table 1.5 has no RHDI column in its index block")
    # The column must sit under the index heading, not a £ heading: the
    # nearest unit heading to its left in the heading row is the index one.
    unit_row = header[index_head[0]]
    units = [
        (j, _text(unit_row[j]))
        for j in range(first, column + 1)
        if j < len(unit_row) and _text(unit_row[j])
    ]
    if not units or not re.match(r"index", units[-1][1], re.I):
        raise EfoTableError("Table 1.5's RHDI column is not an index")

    values = {}
    for row in rows[names_row + 1 :]:
        label = None
        for cell in row[:3]:
            match = QUARTER.match(_text(cell))
            if match:
                label = f"{match[1]}Q{match[2]}"
                break
        if label is None:
            if values:
                break  # the quarterly rows have ended (annual rows follow)
            continue
        value = row[column] if column < len(row) else None
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            if label in values:
                raise EfoTableError(f"Table 1.5 has {label} twice")
            values[label] = float(value)
    if "2008Q1" not in values or abs(values["2008Q1"] - 100) > 1e-6:
        raise EfoTableError("Table 1.5's RHDI index is not 100 in 2008Q1")
    return values
