"""Import HM Treasury's Table 4.1 policy-decisions scorecard.

On Budget day the Treasury publishes Table 4.1, "Policy decisions" (in £
million), as a spreadsheet whose presentational sheet lists each measure's
number, description, head (tax or spend) and Exchequer impact for each
fiscal year, followed by printed totals. This module turns that sheet into
reconciliation source rows, one per measure and fiscal year, and checks the
measure rows against the printed totals, so a mis-keyed figure or a parsing
slip fails loudly instead of flowing into a reconciliation.

Table 4.1 figures are post-behavioural. Static costings come from the policy
costings document and are entered separately (``--static-template`` writes a
blank sheet with the same rows to fill in).

Usage::

    uv run python -m uk_budget_data.treasury_import policy-decisions.xlsx \\
        --source-id AB2026-T4.1 --document "Autumn Budget 2026" \\
        --out data_inputs/treasury/sources/AB2026-T4.1.csv
"""

import argparse
import csv
import math
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

FISCAL_YEAR = re.compile(r"^(\d{4})-(\d{2})$")
# Each printed figure is rounded; the Treasury rounds to the nearest £5m, so
# a total can differ from the sum of its rows by up to £2.5m per row plus
# £2.5m for the total itself. This is a known blind spot: an error smaller
# than the allowance (which grows with the number of rows), or two errors
# that cancel, pass the totals check.
ROUNDING_PER_FIGURE_M = 2.5
NEGLIGIBLE = "*"
# Nil as printed: a hyphen, en, em or figure dash, or a minus sign.
NIL = {"-", "\u2013", "\u2014", "\u2012", "\u2212", ""}

SOURCE_COLUMNS = [
    "source_id",
    "document",
    "table",
    "row_id",
    "row_label",
    "fiscal_year",
    "value_bn",
    "basis",
    "sign_convention",
    "head",
    "notes",
]

TOTAL_LABELS = {
    "total policy decisions": None,
    "total spending policy decisions": "spend",
    "total tax policy decisions": "tax",
}


class ScorecardError(ValueError):
    """Raised when the sheet cannot be parsed unambiguously."""


class TotalsMismatch(ScorecardError):
    """Raised when measure rows do not add up to the printed totals."""


@dataclass
class Measure:
    """One Table 4.1 line."""

    number: str
    label: str
    head: str
    values_m: dict[str, Optional[float]]
    negligible: set[str] = field(default_factory=set)


@dataclass
class Scorecard:
    """A parsed Table 4.1 sheet."""

    sheet: str
    fiscal_years: list[str]
    measures: list[Measure]
    totals_m: dict[Optional[str], dict[str, float]]


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def parse_figure(value) -> tuple[Optional[float], bool]:
    """Parse one cell in £m, returning (value, negligible).

    A nil marker is 0. The negligible marker ``*`` (under £2.5m either way)
    is 0 and flagged. Anything else must be a number.
    """
    if isinstance(value, bool):
        raise ScorecardError(f"unexpected boolean cell {value!r}")
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = _text(value)
        if text in NIL:
            return 0.0, False
        if text == NEGLIGIBLE:
            return 0.0, True
        try:
            number = float(text.replace(",", ""))
        except ValueError as error:
            raise ScorecardError(f"not a figure: {value!r}") from error
    if not math.isfinite(number):
        raise ScorecardError(f"not a finite figure: {value!r}")
    return number, False


# Dashes a fiscal-year label may be printed with: en, em, figure and minus.
_DASHES = str.maketrans(
    {"\u2013": "-", "\u2014": "-", "\u2012": "-", "\u2212": "-"}
)


def _year_label(value) -> str:
    """A fiscal-year header as 2029-30, whichever dash it was printed with."""
    return _text(value).translate(_DASHES)


def _find_header(rows: list[tuple]) -> tuple[int, dict[int, str], int]:
    """Locate the header row, its fiscal-year columns and the head column."""
    for index, row in enumerate(rows):
        years = {
            column: _year_label(cell)
            for column, cell in enumerate(row)
            if FISCAL_YEAR.match(_year_label(cell))
        }
        if len(years) >= 3:
            labels = list(years.values())
            repeated = sorted({y for y in labels if labels.count(y) > 1})
            if repeated:
                raise ScorecardError(
                    f"fiscal year columns repeat: {', '.join(repeated)}"
                )
            head_columns = [
                column
                for column, cell in enumerate(row)
                if _text(cell).lower().startswith("head")
            ]
            if len(head_columns) != 1:
                raise ScorecardError(
                    "the header row needs exactly one 'Head' column"
                )
            stray = [
                _text(cell)
                for column, cell in enumerate(row)
                if column not in years
                and column not in head_columns
                and re.search(r"\d", _text(cell))
            ]
            if stray:
                raise ScorecardError(
                    f"the header row has unrecognised year-like cells: {stray}"
                )
            if head_columns[0] > min(years):
                raise ScorecardError(
                    "the 'Head' column must come before the fiscal years"
                )
            for column, label in years.items():
                start, end = FISCAL_YEAR.match(label).groups()
                if (int(start) + 1) % 100 != int(end):
                    raise ScorecardError(f"malformed fiscal year {label!r}")
            return index, years, head_columns[0]
    raise ScorecardError("no header row with at least three fiscal years")


# Footnote markers that may sit in a cell of their own beside a total's
# label: daggers and superscript figures, at most three characters. A digit
# or an asterisk is not one: in a column whose header was not recognised it
# would be a printed figure or the negligible marker.
_MARKER = re.compile(r"[\u2020\u2021\u00b9\u00b2\u00b3\u2070-\u2079]{1,3}")


def _cells(row: tuple, first_year_column: int) -> list:
    return list(row) + [None] * (first_year_column + 1 - len(row))


def _total_label(label_cells: list[str], row: tuple) -> Optional[str]:
    """The total a row prints, or None if it names no total.

    A total row's labels are one total label and, at most, footnote
    markers. Anything else beside a total label is ambiguous (a heading,
    a note, two totals) and is refused rather than guessed at.
    """
    totals = [cell for cell in label_cells if _total_key(cell) in TOTAL_LABELS]
    if not totals:
        return None
    others = [cell for cell in label_cells if cell not in totals]
    if len(totals) > 1 or any(not _MARKER.fullmatch(cell) for cell in others):
        raise ScorecardError(
            f"a total label beside other text in {row!r}; check this row"
        )
    return _total_key(totals[0])


# A footnote marker printed on the end of a label: superscript figures,
# asterisks, daggers, or up to two digits straight after a letter.
_TRAILING_MARKER = re.compile(
    r"(?:[\u00b9\u00b2\u00b3\u2070-\u2079*\u2020\u2021]+"
    r"|(?<=[a-z])\s*\d{1,2}"
    r"|\s*\((?:\d{1,2}|[a-z])\))$"
)


# Characters that carry no words: format characters (zero-width spaces and
# joiners), and markers or brackets before a label's first letter.
_LEADING_NOISE = re.compile(r"^[^a-z]+")
_TOTAL_WORD = re.compile(r"(?:^|[^a-z])(?:sub-?)?total(?:s)?(?:$|[^a-z])")


def _total_key(label: str) -> str:
    """A label as a total is looked up: lower case, format characters and
    leading markers off, trailing footnote marker off."""
    text = "".join(
        char for char in label if unicodedata.category(char) != "Cf"
    )
    text = _LEADING_NOISE.sub("", text.strip().lower())
    return _TRAILING_MARKER.sub("", text).strip()


def _is_figure(cell) -> bool:
    """Whether a cell holds a number, as a value or as text."""
    if isinstance(cell, bool):
        return False
    if isinstance(cell, (int, float)):
        return True
    return bool(re.fullmatch(r"-?[\d,]+(?:\.\d+)?", _text(cell)))


# A footnote marker printed on the end of a figure: superscript figures,
# asterisks and daggers.
_FIGURE_MARKER = re.compile(
    r"[\u00b9\u00b2\u00b3\u2070-\u2079*\u2020\u2021]+$"
)


# A lettered footnote on a figure: "999 (a)" or "999 [a]" anywhere. A bare
# letter ("999a") is a footnote only where no measure number can be, above
# the header: below it, "12a" is a measure's number.
_BRACKETED_LETTER = re.compile(r"\s*(?:\([a-z]\)|\[[a-z]\])$", re.IGNORECASE)
_BARE_LETTER = re.compile(r"(?<=\d)[a-z]$", re.IGNORECASE)


def _strip_footnotes(text: str, bare_letters: bool) -> str:
    """A figure's text with every footnote marker on its end removed, in
    any order ("999 (a)\u00b9")."""
    while True:
        core = _FIGURE_MARKER.sub("", text).strip()
        core = _BRACKETED_LETTER.sub("", core).strip()
        if bare_letters:
            core = _BARE_LETTER.sub("", core)
        if core == text:
            return core
        text = core


def _figure_like(cell, bare_letters: bool = False) -> bool:
    """Whether a cell could be a printed figure: a number, a nil or
    negligible marker, anything a float can be read from (scientific
    notation included), any of these with footnote markers on the end
    (symbols, superscripts, bracketed letters, and bare letters where
    ``bare_letters``), or digits with no words (a figure in a form the
    importer does not read, such as one printed with a Unicode minus)."""
    if _is_figure(cell):
        return True
    text = _text(cell)
    if text == "" or isinstance(cell, bool):
        return False
    if text in NIL or text == NEGLIGIBLE:
        return True
    core = _strip_footnotes(text, bare_letters)
    if not core:
        # Markers alone: a superscript beside a total's label is no
        # figure, but a leading asterisk is the negligible figure itself,
        # with its own footnote ("*\u2020").
        return text.startswith(NEGLIGIBLE)
    if core in NIL:
        return True
    try:
        float(core.replace(",", "").replace("\u2212", "-"))
        return True
    except ValueError:
        pass
    return bool(re.search(r"\d", core)) and not re.search(r"[^\W\d_]", core)


def _mentions_a_total(label: str) -> bool:
    """Whether a label uses the word total or subtotal anywhere."""
    return bool(_TOTAL_WORD.search(_total_key(label)))


def parse_rows(rows: list[tuple], sheet: str = "") -> Scorecard:
    """Parse a scorecard from its cell values, row by row.

    A total is a row whose labels are a total label and, at most, footnote
    markers, and it must print its figures (the Treasury prints a nil
    figure as "-"). A row that only repeats a total's words with no
    figures, or puts other text beside them, is refused: such a row could
    be a heading, a note or a total, and the importer does not guess.
    """
    header_index, year_columns, head_column = _find_header(rows)
    fiscal_years = [year_columns[c] for c in sorted(year_columns)]
    first_year_column = min(year_columns)
    measures: list[Measure] = []
    totals: dict[Optional[str], dict[str, float]] = {}
    seen_numbers: set[str] = set()

    for row in rows[:header_index]:
        # A title may mention totals; a row that also carries a figure, a
        # nil or a negligible marker could be a printed total and is
        # refused.
        if any(_mentions_a_total(_text(cell)) for cell in row) and any(
            _figure_like(cell, bare_letters=True) for cell in row
        ):
            raise ScorecardError(
                f"a total label above the header in {row!r}; check this row"
            )

    for row in rows[header_index + 1 :]:
        cells = _cells(row, first_year_column)
        raw_values = [
            cells[c] if c < len(cells) else None for c in sorted(year_columns)
        ]
        stray = [
            column
            for column, cell in enumerate(cells)
            if column > head_column
            and column not in year_columns
            and _text(cell)
        ]
        if stray:
            raise ScorecardError(
                f"figures outside the fiscal-year columns in {row!r}"
            )
        label_cells = [_text(c) for c in cells[:head_column] if _text(c)]
        head = _text(cells[head_column]).lower()
        label = max(label_cells, key=len) if label_cells else ""
        blank = all(_text(v) == "" for v in raw_values)
        total = _total_label(label_cells, row)
        # Any other label that begins "Total" is refused, even on a row
        # numbered and headed like a measure: a subtotal with an unfamiliar
        # marker would otherwise be imported as a measure.
        # A blank row with no Head is a heading or footnote ("Totals may not
        # sum due to rounding") and carries no figure, so it is skipped below.
        # A figure left of Head counts too: a subtotal whose only figures
        # sit in a column the header did not name would otherwise look blank
        # and be skipped as a heading.
        carries_figures = (
            not blank
            or head in {"tax", "spend"}
            or any(_figure_like(cell) for cell in cells[:head_column])
        )
        # Total words in the Head column: a total placed one column right,
        # which the blank-row skip below would otherwise drop unread.
        if _mentions_a_total(head):
            raise ScorecardError(
                f"a total label in the Head column in {row!r}; check this row"
            )
        if (
            total is None
            and carries_figures
            and any(_mentions_a_total(cell) for cell in label_cells)
        ):
            raise ScorecardError(
                f"an unrecognised total in {row!r}; check this row"
            )
        if total is not None:
            # A total prints every figure (nil is '-') and has no Head.
            if any(_text(v) == "" for v in raw_values):
                raise ScorecardError(
                    f"{total!r} does not print every year's figure in "
                    f"{row!r} (nil is '-')"
                )
            if _text(cells[head_column]):
                raise ScorecardError(
                    f"{total!r} has text in its Head column in {row!r}; "
                    "check this row"
                )
        is_total = total is not None
        lowered = total if is_total else label.lower()
        if is_total:
            label = total
        if blank and head not in {"tax", "spend"}:
            # A section heading, blank line or footnote. A row with a Tax or
            # Spend head is a measure even when every figure is blank (nil).
            continue
        if is_total:
            if TOTAL_LABELS[lowered] in totals:
                # A second printed total would silently replace the first.
                raise ScorecardError(f"{label!r} appears twice")
            totals[TOTAL_LABELS[lowered]] = {
                year: parse_figure(value)[0]
                for year, value in zip(fiscal_years, raw_values)
            }
            continue
        number_columns = [
            column
            for column, cell in enumerate(cells[:head_column])
            if re.fullmatch(r"\d+[a-z]?", _text(cell))
        ]
        numbers = [_text(cells[column]) for column in number_columns]
        if len(numbers) != 1 or not label or label == numbers[0]:
            raise ScorecardError(
                f"cannot read the measure number and description in {row!r}"
            )
        number = numbers[0]
        # Left of Head a measure prints its number and words. Anything else
        # that could be a figure sits in a column the header did not name
        # as a fiscal year (a blank or unfamiliar header cell), where it
        # would never be read or checked against a total.
        if any(
            _figure_like(cell)
            for column, cell in enumerate(cells[:head_column])
            if column != number_columns[0]
        ):
            raise ScorecardError(
                f"a figure left of the Head column in {row!r}; is a "
                "fiscal-year column missing its header?"
            )
        if number in seen_numbers:
            raise ScorecardError(f"measure {number} appears twice")
        seen_numbers.add(number)
        if head not in {"tax", "spend"}:
            raise ScorecardError(f"measure {number}: head {head!r}")
        values: dict[str, Optional[float]] = {}
        negligible: set[str] = set()
        for year, value in zip(fiscal_years, raw_values):
            figure, is_negligible = parse_figure(value)
            values[year] = figure
            if is_negligible:
                negligible.add(year)
        measures.append(Measure(number, label, head, values, negligible))

    if not measures:
        raise ScorecardError("no measure rows found")
    return Scorecard(sheet, fiscal_years, measures, totals)


def check_totals(scorecard: Scorecard) -> list[str]:
    """Compare measure rows with every printed total.

    Returns one line per total checked. Raises TotalsMismatch, listing each
    year that differs by more than the rounding allowance.
    """
    if None not in scorecard.totals_m:
        raise TotalsMismatch("the sheet prints no 'Total policy decisions'")
    report = []
    problems = []
    for head, printed in scorecard.totals_m.items():
        rows = [
            m for m in scorecard.measures if head is None or m.head == head
        ]
        allowance = ROUNDING_PER_FIGURE_M * (len(rows) + 1)
        for year in scorecard.fiscal_years:
            computed = sum(m.values_m[year] for m in rows)
            gap = computed - printed[year]
            label = head or "all"
            if abs(gap) > allowance:
                problems.append(
                    f"{label} {year}: rows sum to {computed:,.0f}m, "
                    f"printed total {printed[year]:,.0f}m"
                )
            report.append(
                f"{label} {year}: rows {computed:,.0f}m vs printed "
                f"{printed[year]:,.0f}m (allowance ±{allowance:g}m)"
            )
    if problems:
        raise TotalsMismatch("; ".join(problems))
    return report


def read_workbook(path: Path, sheet: Optional[str] = None) -> Scorecard:
    """Read Table 4.1 from a workbook (cached values, not formulas)."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    if sheet is None:
        candidates = [
            ws.title
            for ws in workbook.worksheets
            if "scorecard" in ws.title.lower()
        ]
        sheet = (
            candidates[0] if len(candidates) == 1 else workbook.sheetnames[0]
        )
    worksheet = workbook[sheet]
    # Read every cell, not the used range the file declares: an exporter
    # that writes a stale range would otherwise hide a year's column.
    worksheet.reset_dimensions()
    rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
    return parse_rows(rows, sheet)


def source_rows(
    scorecard: Scorecard,
    source_id: str,
    document: str,
    table: str = "4.1",
    basis: str = "post_behavioural",
) -> list[dict]:
    """Reconciliation source rows: one per measure and fiscal year."""
    rows = []
    for measure in scorecard.measures:
        for year in scorecard.fiscal_years:
            rows.append(
                {
                    "source_id": source_id,
                    "document": document,
                    "table": table,
                    "row_id": measure.number,
                    "row_label": measure.label,
                    "fiscal_year": year,
                    "value_bn": f"{measure.values_m[year] / 1000:.4f}",
                    "basis": basis,
                    "sign_convention": "exchequer_positive",
                    "head": measure.head,
                    "notes": (
                        "negligible (*)" if year in measure.negligible else ""
                    ),
                }
            )
    return rows


def static_template(rows: list[dict], source_id: str) -> list[dict]:
    """Blank static-costing rows matching the Table 4.1 rows."""
    return [
        {
            **row,
            "source_id": source_id,
            "table": "policy costings",
            "value_bn": "",
            "basis": "static_costing",
            "notes": "fill from the policy costings document",
        }
        for row in rows
    ]


def write_csv(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=SOURCE_COLUMNS, lineterminator="\n"
        )
        writer.writeheader()
        writer.writerows(rows)


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m uk_budget_data.treasury_import",
        description="Import HM Treasury's Table 4.1 policy decisions.",
    )
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--source-id", required=True)
    parser.add_argument("--document", required=True)
    parser.add_argument("--sheet", default=None)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument(
        "--static-template",
        type=Path,
        default=None,
        help="also write a blank static-costing sheet with the same rows",
    )
    args = parser.parse_args(argv)
    try:
        scorecard = read_workbook(args.workbook, args.sheet)
        report = check_totals(scorecard)
    except ScorecardError as error:
        print(f"error: {error}", file=sys.stderr)
        return 1
    rows = source_rows(scorecard, args.source_id, args.document)
    write_csv(rows, args.out)
    if args.static_template:
        write_csv(
            static_template(rows, f"{args.source_id}-static"),
            args.static_template,
        )
    print(
        f"{len(scorecard.measures)} measures × "
        f"{len(scorecard.fiscal_years)} years from sheet "
        f"{scorecard.sheet!r} -> {args.out}"
    )
    for line in report:
        print(line)
    return 0


if __name__ == "__main__":
    sys.exit(main())
