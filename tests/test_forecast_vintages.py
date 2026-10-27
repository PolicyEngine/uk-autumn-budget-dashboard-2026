"""The Forecast vintages build (src/uk_budget_data/forecast_vintages/).

The committed file must be exactly what the build makes from the committed
inputs, so it is never edited by hand. The build refuses inputs whose bytes
differ from their pinned SHA-256, outturn that strays from the OBR's own,
forecasts with gaps and forecasts after March 2026. The readers are tested
on small synthetic workbooks and series in the publishers' layouts, and the
period rules by properties over every month.
"""

import dataclasses
import decimal
import hashlib
import json
import math
import shutil
from pathlib import Path

import httpx
import openpyxl
import pytest
from hypothesis import given, settings
from hypothesis import strategies as st

from uk_budget_data.cli import parse_args
from uk_budget_data.forecast_vintages import build as B
from uk_budget_data.forecast_vintages import (
    efo,
    fetch,
    hofd,
    ons,
    periods,
    sources,
)
from uk_budget_data.forecast_vintages import variables as V

pytestmark = pytest.mark.xdist_group("forecast_vintages")

# No property here is about time, and on a busy host an example that takes
# under a millisecond has taken 295 ms (review round 19), so none of them
# has Hypothesis's 200 ms deadline
NO_DEADLINE = settings(deadline=None)

ROOT = Path(__file__).resolve().parent.parent
COMMITTED = ROOT / B.OUTPUT_PATH
REVIEW_CSV = ROOT / "data_inputs" / "public_data_review.csv"


@pytest.fixture(scope="module")
def built():
    return B.build(ROOT)


@pytest.fixture(scope="module")
def committed():
    return json.loads(COMMITTED.read_text())


# --- The committed file -----------------------------------------------------


def test_the_committed_file_is_exactly_what_the_build_makes(built):
    assert B.output_text(built) == COMMITTED.read_text()


def test_every_pinned_obr_file_matches_its_sha256():
    for file in (sources.HOFD,) + sources.EFO_TABLES:
        data = (ROOT / file.path).read_bytes()
        assert hashlib.sha256(data).hexdigest() == file.sha256, file.id


def test_the_inputs_readme_lists_every_pinned_file():
    readme = (ROOT / sources.INPUTS_DIR / "README.md").read_text()
    for file in (sources.HOFD,) + sources.EFO_TABLES:
        assert file.sha256 in readme, file.id
        if file.timestamp:
            assert file.archive_url in readme, file.id
    for cdid in sources.ONS:
        assert f"[{cdid}]" in readme, cdid


def test_the_review_row_is_pending_on_the_committed_bytes():
    digest = hashlib.sha256(COMMITTED.read_bytes()).hexdigest()
    rows = [
        line.split(",")
        for line in REVIEW_CSV.read_text().splitlines()
        if line.startswith(B.OUTPUT_PATH.as_posix() + ",")
    ]
    assert len(rows) == 1
    path, kind, sha, reviewer = rows[0][:4]
    assert (kind, sha) == ("pending", digest)
    assert reviewer


def test_the_file_names_every_source_with_its_sha256(committed):
    meta = committed["meta"]["sources"]
    assert [s["id"] for s in meta["obr"]] == [
        f.id for f in (sources.HOFD,) + sources.EFO_TABLES
    ]
    assert {s["cdid"] for s in meta["ons"]} == set(sources.ONS)
    for source in meta["obr"] + meta["ons"] + [meta["hpi"]]:
        assert len(source["sha256"]) == 64
    for source in meta["obr"]:
        if source["archive_url"]:
            assert source["archive_url"].startswith(
                "https://web.archive.org/web/"
            )
            assert source["archive_url"].endswith(source["url"])


def test_a_sheet_title_ending_with_its_unit_names_it_once():
    sheet = hofd.Sheet(
        name="£PSNB",
        title="Public sector net borrowing (£ billion)",
        unit="£ billion",
        notes=(),
        periods=(),
        vintages=(),
        outturn={},
    )
    assert B._sheet_source(sheet).endswith("borrowing (£ billion)")
    bare = dataclasses.replace(sheet, title="Public sector net borrowing")
    assert B._sheet_source(bare).endswith("borrowing (£ billion)")


def test_every_variable_has_its_units_measures_and_forecasts(committed):
    ids = [v["id"] for v in committed["variables"]]
    assert ids == [v.id for v in V.VARIABLES]
    for variable in committed["variables"]:
        assert variable["default_unit"] in variable["units"]
        for unit, spec in variable["units"].items():
            assert unit in ("level", "index", "pct")
            assert variable[spec["from"]] is not None
        for name in ("level", "share"):
            measure = variable[name]
            if measure is None:
                continue
            assert measure["kind"] in ("growth", "index", "level")
            assert measure["outturn"]
            values = list(measure["outturn"].values()) + [
                x for v in measure["vintages"] for x in v["values"].values()
            ]
            assert all(
                isinstance(x, float) and math.isfinite(x) for x in values
            )


def test_forecasts_run_june_2010_to_march_2026_in_order(committed):
    for variable in committed["variables"]:
        for measure in filter(None, (variable["level"], variable["share"])):
            dates = [(v["year"], v["month"]) for v in measure["vintages"]]
            assert dates == sorted(dates)
            assert len(set(dates)) == len(dates)
            assert dates[-1] == B.NEWEST_VINTAGE
            if variable["frequency"] != "quarterly":
                assert dates[0] == hofd.FIRST_OBR_VINTAGE
                assert len(dates) == 33


def test_every_forecast_starts_at_its_launch_and_has_no_gaps(committed):
    for variable in committed["variables"]:
        frequency = variable["frequency"]
        for measure in filter(None, (variable["level"], variable["share"])):
            for vintage in measure["vintages"]:
                launch = periods.launch(
                    frequency, vintage["year"], vintage["month"]
                )
                assert vintage["launch"] == str(launch)
                keys = list(vintage["values"])
                first = (
                    launch
                    if measure["kind"] == "index"
                    else periods.next_period(launch)
                )
                allowed = {str(launch), str(first)}
                assert keys[0] in allowed, (variable["id"], vintage["label"])
                assert str(first) in keys
                parsed = [_period(k, frequency) for k in keys]
                for a, b in zip(parsed, parsed[1:]):
                    assert periods.next_period(a) == b


def _period(key: str, frequency: str):
    return int(key) if frequency == "calendar" else key


def test_the_outturn_matches_the_obrs_within_tolerance(committed):
    for variable in committed["variables"]:
        for measure in filter(None, (variable["level"], variable["share"])):
            check = measure["check"]
            assert len(check["periods"]) == B.CHECK_PERIODS
            if check["tolerance_is_share"]:
                continue  # the share is checked by the build; the gap is in units
            assert abs(check["largest_gap"]) <= check["tolerance"]


def test_quarterly_income_per_head_matches_the_published_chart(committed):
    # Max's chart (25 October): £25,820 at the start of the parliament,
    # £26,308 latest, £26,628 at its end on the March 2026 forecast carried
    # on from the latest outturn, all quarterly annualised (x4).
    variable = next(
        v for v in committed["variables"] if v["id"] == "rhdi_per_person"
    )
    outturn = variable["level"]["outturn"]
    newest = variable["level"]["vintages"][-1]["values"]
    start, end = (
        committed["meta"]["parliament"][k]["quarterly"]
        for k in ("start", "end")
    )
    latest = list(outturn)[-1]
    assert round(outturn[start] * 4) == 25820
    assert (latest, round(outturn[latest] * 4)) == ("2026Q2", 26308)
    assert round(outturn[latest] * newest[end] / newest[latest] * 4) == 26628


def test_no_policyengine_figures(committed):
    assert "policyengine" not in COMMITTED.read_text().lower()


# --- Refusals ---------------------------------------------------------------


def test_a_changed_obr_file_is_refused(tmp_path):
    target = tmp_path / sources.HOFD.path
    target.parent.mkdir(parents=True)
    data = bytearray((ROOT / sources.HOFD.path).read_bytes())
    data[-1] ^= 1
    target.write_bytes(bytes(data))
    with pytest.raises(B.BuildError, match="sha256"):
        B.verify_archived(sources.HOFD, tmp_path)


def test_a_missing_obr_file_says_how_to_fetch_it(tmp_path):
    with pytest.raises(B.BuildError, match="--fetch"):
        B.verify_archived(sources.HOFD, tmp_path)


def test_an_ons_file_that_differs_from_its_record_is_refused(tmp_path):
    shutil.copytree(
        ROOT / sources.INPUTS_DIR / "ons",
        tmp_path / sources.INPUTS_DIR / "ons",
    )
    retrieved = json.loads((ROOT / sources.RETRIEVED_PATH).read_text())
    retrieved["ons"]["CRXX"]["sha256"] = "0" * 64
    with pytest.raises(B.BuildError, match="CRXX|crxx"):
        B.load_ons(tmp_path, retrieved)


def test_an_unrecorded_ons_series_is_refused(tmp_path):
    retrieved = json.loads((ROOT / sources.RETRIEVED_PATH).read_text())
    del retrieved["ons"]["ABMI"]
    with pytest.raises(B.BuildError, match="ABMI"):
        B.load_ons(ROOT, retrieved)


def _measure(**kwargs):
    base = dict(
        unit="£ billion",
        kind="level",
        outturn=lambda s, h: {},
        outturn_source="",
        sheet="X",
    )
    return V.Measure(**{**base, **kwargs})


def _variable(frequency="fiscal"):
    return dataclasses.replace(V.VARIABLES[0], id="test", frequency=frequency)


def test_outturn_beyond_tolerance_is_refused():
    row = {periods.fiscal_label(y): 100.0 for y in range(2010, 2020)}
    near = {k: v + 0.9 for k, v in row.items()}
    far = dict(near, **{"2018-19": 102.0})
    measure = _measure(tolerance=1.0)
    check = B.check_outturn(_variable(), measure, near, row)
    assert check["largest_gap"] == pytest.approx(0.9)
    with pytest.raises(B.BuildError, match="2018-19"):
        B.check_outturn(_variable(), measure, far, row)


def test_a_relative_tolerance_scales_with_the_value():
    row = {periods.fiscal_label(y): 1000.0 for y in range(2010, 2020)}
    measure = _measure(tolerance=0.02, relative=True)
    B.check_outturn(_variable(), measure, {k: 1019.0 for k in row}, row)
    with pytest.raises(B.BuildError):
        B.check_outturn(_variable(), measure, {k: 1021.0 for k in row}, row)


def test_growth_outturn_is_compared_as_growth():
    levels = {y: 100.0 * 1.02 ** (y - 2000) for y in range(2000, 2020)}
    row = {y: 2.0 for y in range(2001, 2020)}
    check = B.check_outturn(
        _variable("calendar"), _measure(kind="growth"), levels, row
    )
    assert check["largest_gap"] == pytest.approx(0.0, abs=1e-9)
    assert check["in"] == "% change on a year earlier"


def test_too_few_periods_to_check_is_refused():
    row = {2020: 1.0, 2021: 1.0}
    with pytest.raises(B.BuildError, match="only 2"):
        B.check_outturn(_variable("calendar"), _measure(), row, row)


def test_a_forecast_with_a_gap_is_refused():
    values = {2024: 1.0, 2025: 1.0, 2027: 1.0}
    with pytest.raises(B.BuildError, match="skips"):
        B.vintage_record(
            _variable("calendar"), _measure(), "March 2025", 2025, 3, values
        )


def test_a_forecast_without_its_first_period_is_refused():
    with pytest.raises(B.BuildError, match="no value for 2026"):
        B.vintage_record(
            _variable("calendar"),
            _measure(kind="growth"),
            "March 2026",
            2026,
            3,
            {2025: 1.0, 2027: 1.0},
        )


def test_a_forecast_keeps_its_values_from_launch_on():
    values = {y: float(y) for y in range(2020, 2031)}
    record = B.vintage_record(
        _variable("calendar"), _measure(), "March 2026", 2026, 3, values
    )
    assert record["launch"] == "2025"
    assert list(record["values"]) == [str(y) for y in range(2025, 2031)]
    assert record["id"] == "march_2026"


def test_a_forecast_after_march_2026_is_refused():
    def vintage(year, month):
        return hofd.Vintage(label="x", year=year, month=month, values={})

    with pytest.raises(B.BuildError, match="October 2026|post-Budget"):
        B._check_newest(
            "x", [vintage(2026, 3), hofd.Vintage("October 2026", 2026, 10, {})]
        )
    with pytest.raises(B.BuildError, match="not March 2026"):
        B._check_newest("x", [vintage(2025, 11)])
    B._check_newest("x", [vintage(2025, 11), vintage(2026, 3)])


# --- Readers ----------------------------------------------------------------


def _hofd_workbook(path, rows, periods_row):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "UKGDP"
    ws.append(["Back to contents", "Real GDP growth"])
    ws.append(["Per cent"])
    ws.append([None])
    ws.append(["Back to contents"] + periods_row)
    for row in rows:
        ws.append(row)
    wb.save(path)
    return hofd.open_workbook(path)


def test_the_database_reader_keeps_the_obrs_forecasts_only(tmp_path):
    workbook = _hofd_workbook(
        tmp_path / "hofd.xlsx",
        [
            ["July 1996", 1.0, 2.0, None],
            ["July 1996", 1.5, 2.5, None],  # HM Treasury's rows repeat names
            ["June 2010", 1.2, None, "n/a"],
            ["March2026", None, 1.1, 1.4],  # the database's own typo
            ["Outturn data*", 1.3, 1.9, None],
            ["Memo: March 2019 restated", 9.0, 9.0, 9.0],
            ["* Outturn as available at last forecast"],
        ],
        [2009, 2010, 2011],
    )
    sheet = hofd.read_sheet(workbook, "UKGDP")
    assert [v.label for v in sheet.vintages] == ["June 2010", "March 2026"]
    # A blank or text cell is missing, never zero
    assert sheet.vintages[0].values == {2009: 1.2}
    assert sheet.vintages[1].values == {2010: 1.1, 2011: 1.4}
    assert sheet.outturn == {2009: 1.3, 2010: 1.9}
    assert sheet.periods == (2009, 2010, 2011)


def test_the_database_reader_reads_fiscal_years(tmp_path):
    workbook = _hofd_workbook(
        tmp_path / "hofd.xlsx",
        [["March 2026", 1.0, 2.0], ["Outturn data", 1.0, None]],
        ["2024-25", "2025-26"],
    )
    assert hofd.read_sheet(workbook, "UKGDP").vintages[0].values == {
        "2024-25": 1.0,
        "2025-26": 2.0,
    }


def test_the_database_reader_refuses_forecasts_out_of_order(tmp_path):
    workbook = _hofd_workbook(
        tmp_path / "hofd.xlsx",
        [["March 2026", 1.0], ["November 2025", 1.0], ["Outturn data", 1.0]],
        [2025],
    )
    with pytest.raises(hofd.HofdError, match="date order"):
        hofd.read_sheet(workbook, "UKGDP")


def test_the_database_reader_needs_an_outturn_row(tmp_path):
    workbook = _hofd_workbook(
        tmp_path / "hofd.xlsx", [["March 2026", 1.0]], [2025]
    )
    with pytest.raises(hofd.HofdError, match="outturn"):
        hofd.read_sheet(workbook, "UKGDP")


@NO_DEADLINE
@given(
    st.integers(1990, 2100),
    st.sampled_from(hofd.MONTH_NAMES),
    st.sampled_from(["", " ", "  "]),
)
def test_vintage_labels_parse_to_their_month(year, month, space):
    assert hofd.parse_vintage_label(f"{month}{space}{year}") == (
        year,
        hofd.MONTH_NAMES.index(month) + 1,
    )


def test_labels_that_are_not_months_are_not_forecasts():
    for text in (
        "Outturn data",
        "Memo: March 2019",
        "Marching 2020",
        None,
        "2020",
    ):
        assert hofd.parse_vintage_label(text) is None


def _table_15(path, rows, heading="Real household disposable income"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1.5"
    ws.append([None, "Table 1.5: Per person"])
    ws.append([None, None, "By total population", None, "Ages 16+"])
    ws.append([None, None, "Index: 2008Q1=100", None, "Index: 2008Q1=100"])
    ws.append([None, None, heading, None, heading])
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def test_table_15_gives_the_total_population_index(tmp_path):
    path = _table_15(
        tmp_path / "t.xlsx",
        [
            [None, "2008Q1", 100.0, None, 100.0],
            [None, "2008 Q2", 101.5, None, 99.0],
            [None, "2008-09", 103.0, None, 98.0],  # annual rows end the run
        ],
    )
    assert efo.read_rhdi_per_person(path) == {"2008Q1": 100.0, "2008Q2": 101.5}


def test_table_15_accepts_a_footnote_marker(tmp_path):
    path = _table_15(
        tmp_path / "t.xlsx",
        [[None, "2008Q1", 100.0, None, 100.0]],
        heading="Real household disposable income1",
    )
    assert efo.read_rhdi_per_person(path) == {"2008Q1": 100.0}


def test_table_15_must_be_indexed_to_2008q1(tmp_path):
    path = _table_15(
        tmp_path / "t.xlsx", [[None, "2008Q1", 98.0, None, 100.0]]
    )
    with pytest.raises(efo.EfoTableError, match="100 in 2008Q1"):
        efo.read_rhdi_per_person(path)


def test_a_workbook_without_table_15_is_refused(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "1.4"
    wb.save(tmp_path / "t.xlsx")
    with pytest.raises(efo.EfoTableError, match="Table 1.5"):
        efo.read_rhdi_per_person(tmp_path / "t.xlsx")


def test_every_committed_table_15_reads(committed):
    variable = next(
        v for v in committed["variables"] if v["id"] == "rhdi_per_person"
    )
    labels = [v["label"] for v in variable["level"]["vintages"]]
    assert labels == [f.vintage for f in sources.EFO_TABLES]


def _ons_file(
    path, cdid="CRXX", dataset="UKEA", quarters=(), years=(), months=()
):
    doc = {
        "description": {
            "cdid": cdid,
            "datasetId": dataset,
            "title": "A series",
            "unit": "£",
            "releaseDate": "2026-09-29T23:00:00.000Z",
        },
        "years": [{"date": d, "value": v} for d, v in years],
        "quarters": [{"date": d, "value": v} for d, v in quarters],
        "months": [{"date": d, "value": v} for d, v in months],
    }
    path.write_text(json.dumps(doc))
    return path


def test_an_ons_file_of_another_series_is_refused(tmp_path):
    path = _ons_file(tmp_path / "x.json", cdid="ABMI")
    with pytest.raises(ons.OnsSeriesError, match="ABMI"):
        ons.read_series(path, "CRXX", "UKEA")
    with pytest.raises(ons.OnsSeriesError, match="dataset"):
        ons.read_series(_ons_file(tmp_path / "y.json"), "CRXX", "QNA")


def test_blank_ons_values_are_missing_not_zero(tmp_path):
    path = _ons_file(
        tmp_path / "x.json",
        quarters=[("2025 Q4", "6500"), ("2026 Q1", ""), ("2026 Q2", " ")],
        years=[("2025", "26000")],
        months=[("2026 MAR", "95.1")],
    )
    series = ons.read_series(path, "CRXX", "UKEA")
    assert series.quarters == {"2025Q4": 6500.0}
    assert series.years == {2025: 26000.0}
    assert series.months == {"2026-03": 95.1}
    assert series.release_date == "2026-09-29"


@NO_DEADLINE
@given(
    st.dictionaries(st.integers(2000, 2010), st.floats(-1e6, 1e6), min_size=1)
)
def test_fiscal_years_sum_complete_runs_of_four_quarters(per_year):
    quarters = {}
    for year, value in per_year.items():
        for q in range(1, 5):
            quarters[f"{year}Q{q}"] = value
    totals = ons.fiscal_sum(quarters)
    for label, total in totals.items():
        start = int(label[:4])
        assert start in per_year and start + 1 in per_year
        assert total == pytest.approx(
            3 * per_year[start] + per_year[start + 1]
        )
    for start in per_year:
        if start + 1 in per_year:
            assert periods.fiscal_label(start) in totals


def test_debt_is_taken_at_the_end_of_march():
    assert ons.fiscal_end_march(
        {"2025-03": 95.0, "2025-04": 96.0, "2026-03": 94.0}
    ) == {
        "2024-25": 95.0,
        "2025-26": 94.0,
    }


def test_house_prices_average_complete_years_only(tmp_path):
    months = {f"2024-{m:02d}": 100.0 * m for m in range(1, 13)}
    months["2025-01"] = 5.0
    by_year = V._house_prices({}, months)
    assert by_year == {2024: pytest.approx(650.0)}


# --- Values that are not numbers, and growth across a gap -------------------


@pytest.mark.parametrize("text", ["NaN", "nan", "inf", "-Infinity"])
def test_an_ons_value_that_is_not_a_number_is_refused(tmp_path, text):
    path = _ons_file(
        tmp_path / "x.json", years=[("2024", "1.0"), ("2025", text)]
    )
    with pytest.raises(ons.OnsSeriesError, match="finite"):
        ons.read_series(path, "CRXX", "UKEA")


@pytest.mark.parametrize("value", [math.nan, math.inf, -math.inf])
def test_an_obr_value_that_is_not_a_number_is_refused(value):
    with pytest.raises(hofd.HofdError, match="finite"):
        hofd._number(value)


def test_a_house_price_that_is_not_a_number_is_refused(tmp_path):
    path = tmp_path / "hpi.csv"
    path.write_text(
        '"Name","URI","Period","Average price All property types"\n'
        '"United Kingdom","x","2024-01","290000"\n'
        '"United Kingdom","x","2024-02","NaN"\n'
    )
    with pytest.raises(B.BuildError, match="2024-02"):
        B.read_hpi(path)


def test_an_outturn_that_is_not_a_number_is_refused():
    row = {periods.fiscal_label(y): 100.0 for y in range(2010, 2020)}
    mine = dict(row, **{"2015-16": math.nan})
    with pytest.raises(B.BuildError, match="2015-16"):
        B.check_outturn(_variable(), _measure(tolerance=1.0), mine, row)


def test_the_file_is_never_written_with_a_value_that_is_not_a_number():
    with pytest.raises(ValueError):
        B.output_text({"values": {"2025": math.nan}})


def test_growth_needs_the_period_before():
    assert ons.growth({2020: 100.0, 2022: 121.0}) == {}
    assert ons.growth({2020: 100.0, 2021: 110.0, 2023: 121.0}) == {
        2021: pytest.approx(10.0)
    }
    assert ons.growth({"2023-24": 100.0, "2025-26": 110.0}) == {}
    assert ons.growth({"2024Q4": 100.0, "2025Q1": 101.0}) == {
        "2025Q1": pytest.approx(1.0)
    }


@NO_DEADLINE
@given(st.sets(st.integers(1950, 2030), min_size=1, max_size=40))
def test_growth_is_only_ever_on_the_year_before(years):
    levels = {y: 100.0 + y for y in years}
    out = ons.growth(levels)
    assert set(out) == {y for y in years if y - 1 in years}
    for y, rate in out.items():
        assert rate == pytest.approx((levels[y] / levels[y - 1] - 1) * 100)
    assert list(out) == sorted(out)


def test_growth_outturn_two_years_apart_has_nothing_to_check():
    # Every other year, each 2% up: no year has the year before it
    levels = {
        y: 100.0 * 1.02 ** ((y - 2000) // 2) for y in range(2000, 2018, 2)
    }
    row = {y: 2.0 for y in range(2010, 2018)}
    with pytest.raises(B.BuildError, match="only 0"):
        B.check_outturn(
            _variable("calendar"), _measure(kind="growth"), levels, row
        )


# --- A forecast's figures read as its kind ------------------------------------

KIND_CASES = json.loads(
    (
        ROOT / "tests" / "fixtures" / "forecast_vintages_kind_cases.json"
    ).read_text()
)["cases"]


def _apply_kind_op(data: dict, op: dict) -> None:
    """One change from the shared kind cases, made in place."""
    variable = next(v for v in data["variables"] if v["id"] == op["variable"])
    measure = variable[op["measure"]]
    frequency = variable["frequency"]
    lag = 4 if frequency == "quarterly" else 1

    def no(key):
        return periods.sort_key(B._period(key, frequency))

    vintage = next(
        (v for v in measure["vintages"] if v["label"] == op.get("vintage")),
        None,
    )
    if op["op"] == "kind":
        measure["kind"] = op["to"]
    elif op["op"] == "drop_launch":
        del vintage["values"][str(vintage["launch"])]
    elif op["op"] == "drop_every_launch":
        for v in measure["vintages"]:
            if v["label"] != op.get("but"):
                v["values"].pop(str(v["launch"]), None)
    elif op["op"] == "set_after_launch":
        for key in vintage["values"]:
            if key != str(vintage["launch"]):
                vintage["values"][key] = op["to"]
    elif op["op"] == "set_last":
        vintage["values"][max(vintage["values"], key=no)] = op["to"]
    elif op["op"] == "scale":
        for key in vintage["values"]:
            vintage["values"][key] *= op["by"]
    elif op["op"] == "launch_off_rate":
        launch = str(vintage["launch"])
        before = next(
            x
            for k, x in measure["outturn"].items()
            if no(k) == no(launch) - lag
        )
        rate = 100 * (measure["outturn"][launch] / before - 1)
        vintage["values"][launch] = rate + op["by"]
    elif op["op"] == "launch_off_level":
        launch = str(vintage["launch"])
        size = max(abs(x) for x in measure["outturn"].values())
        vintage["values"][launch] = (
            measure["outturn"][launch] + op["by"] * size
        )
    elif op["op"] == "keep_outturn":
        measure["outturn"] = {
            k: x for k, x in measure["outturn"].items() if k in op["periods"]
        }
    else:
        raise ValueError(f"Unknown kind case op {op['op']!r}")


def _kind_check(data: dict, variable_id: str, name: str) -> None:
    variable = next(v for v in data["variables"] if v["id"] == variable_id)
    measure = variable[name]
    B.check_kind(
        variable_id,
        variable["frequency"],
        measure["kind"],
        measure["outturn"],
        measure["vintages"],
    )


@pytest.mark.parametrize("case", KIND_CASES, ids=lambda c: c["name"])
def test_the_build_reads_kinds_as_the_page_does(committed, case):
    # The page's loader (kindProblem) is held to the same cases in
    # src/utils/forecastVintages/kind.test.js, so the two agree
    data = json.loads(json.dumps(committed))
    for op in case["ops"]:
        _apply_kind_op(data, op)
    target = case["ops"][0]
    if case["accepted"]:
        _kind_check(data, target["variable"], target["measure"])
    else:
        with pytest.raises(B.BuildError):
            _kind_check(data, target["variable"], target["measure"])


def test_every_committed_measure_reads_as_its_kind(committed):
    for variable in committed["variables"]:
        for name in ("level", "share"):
            if variable.get(name):
                _kind_check(committed, variable["id"], name)


def _committed_measures():
    data = json.loads(COMMITTED.read_text())
    return [
        (v["id"], name)
        for v in data["variables"]
        for name in ("level", "share")
        if v.get(name)
    ]


@NO_DEADLINE
@given(
    st.sampled_from(_committed_measures()),
    st.integers(0, 32),
    st.floats(0.01, 99.99) | st.floats(-99.99, -0.01),
)
def test_forecasts_after_the_launch_are_never_set_against_outturn(
    target, which, value
):
    # A forecast may miss by any amount: whatever its figures after its
    # launch period (an index's above zero), the build still writes it
    data = json.loads(COMMITTED.read_text())
    variable_id, name = target
    variable = next(v for v in data["variables"] if v["id"] == variable_id)
    measure = variable[name]
    vintage = measure["vintages"][which % len(measure["vintages"])]
    figure = abs(value) if measure["kind"] == "index" else value
    if measure["kind"] == "level":
        figure = value * 1e6
    for key in vintage["values"]:
        if key != str(vintage["launch"]):
            vintage["values"][key] = figure
    _kind_check(data, variable_id, name)


def test_a_share_tolerance_needs_a_forecast_of_levels():
    levels = {y: 100.0 * 1.02 ** (y - 2000) for y in range(2000, 2020)}
    row = {y: 2.0 for y in range(2001, 2020)}
    with pytest.raises(B.BuildError, match="share"):
        B.check_outturn(
            _variable("calendar"),
            _measure(kind="growth", relative=True, tolerance=0.02),
            levels,
            row,
        )


def test_a_tolerance_finer_than_a_gap_is_written_is_refused():
    row = {periods.fiscal_label(y): 100.0 for y in range(2010, 2020)}
    with pytest.raises(B.BuildError, match="decimals"):
        B.check_outturn(_variable(), _measure(tolerance=0.00001), row, row)


def test_an_index_is_never_checked_against_a_database_sheet():
    # It is checked in growth rates against Table 1.5 (check_index); set
    # against a sheet as it stands, its gaps would be in index points
    row = {periods.fiscal_label(y): 100.0 for y in range(2010, 2020)}
    with pytest.raises(B.BuildError, match="Table 1.5"):
        B.check_outturn(_variable(), _measure(kind="index"), row, row)


def test_a_check_is_written_in_its_kinds_terms():
    levels = {y: 100.0 * 1.02 ** (y - 2000) for y in range(2000, 2020)}
    rates = {y: 2.0 for y in range(2001, 2020)}
    growth = B.check_outturn(
        _variable("calendar"), _measure(kind="growth"), levels, rates
    )
    assert growth["in"] == "% change on a year earlier"
    level = B.check_outturn(
        _variable("calendar"), _measure(unit="£ billion"), levels, levels
    )
    assert level["in"] == "£ billion"


def test_a_gap_within_its_tolerance_only_before_rounding_is_refused():
    # 0.00026 is within a tolerance of 0.00026, but is written 0.0003: the
    # page would state a gap beyond the tolerance it states
    row = {y: 100.0 for y in range(2010, 2020)}
    mine = {k: v + 0.00026 for k, v in row.items()}
    variable, measure = _variable("calendar"), _measure(tolerance=0.00026)
    check = B.check_outturn(variable, measure, mine, row)
    assert check["largest_gap"] > measure.tolerance
    with pytest.raises(B.BuildError, match="beyond its tolerance"):
        B.check_as_read(
            "test", "calendar", "level", "£ billion", check, B._ordered(mine)
        )


def test_a_mistyped_tolerance_is_refused_as_the_page_reads_it():
    # check_outturn takes True for 1; the page refuses a tolerance that is
    # not a number, so the build does too
    row = {y: 100.0 for y in range(2010, 2020)}
    check = B.check_outturn(
        _variable("calendar"), _measure(tolerance=True), row, row
    )
    with pytest.raises(B.BuildError, match="incomplete"):
        B.check_as_read(
            "test",
            "calendar",
            "level",
            "£ billion",
            json.loads(json.dumps(check)),
            B._ordered(row),
        )


def _units(variable_id: str, unit: str, **spec):
    variable = next(v for v in V.VARIABLES if v.id == variable_id)
    units = {**variable.units, unit: {**variable.units[unit], **spec}}
    return dataclasses.replace(variable, units=units)


def test_every_variables_units_can_be_drawn():
    for variable in V.VARIABLES:
        B.check_units(variable)


@pytest.mark.parametrize(
    "variable",
    [
        _units("rhdi_per_person", "level", annualise=False),
        dataclasses.replace(
            _units("rhdi_per_person", "level"),
            units={
                **_units("rhdi_per_person", "level").units,
                "level": {
                    k: x
                    for k, x in _units("rhdi_per_person", "level")
                    .units["level"]
                    .items()
                    if k != "annualise"
                },
            },
        ),
        _units("real_gdp", "level", annualise=True),
        _units("rhdi_per_person", "index", annualise=True),
    ],
    ids=["quarterly £ not annual", "quarterly £ unsaid", "annual £", "index"],
)
def test_a_quarterly_measures_pounds_and_only_those_are_at_an_annual_rate(
    variable,
):
    # Its axis names a year and its figures are a quarter's; the page's
    # loader refuses a file otherwise
    with pytest.raises(B.BuildError, match="annual rate"):
        B.check_units(variable)


def _without_transform(variable_id: str, unit: str):
    variable = _units(variable_id, unit)
    spec = {k: x for k, x in variable.units[unit].items() if k != "transform"}
    return dataclasses.replace(variable, units={**variable.units, unit: spec})


@pytest.mark.parametrize(
    "variable, message",
    [
        # Found by the eleventh review: a rate's change on a year earlier,
        # under the rate's own axis
        (_units("psnd", "pct", transform="growth"), "other units"),
        (_units("unemployment", "pct", transform="growth"), "other units"),
        (_units("psnb", "pct", transform="growth"), "taken of no level"),
        (_without_transform("real_gdp", "pct"), "other units"),
        (_units("psnb", "pct", **{"from": "level"}), "other units"),
        (_units("psnb", "level", **{"from": "share"}), "other units"),
        (_units("real_gdp", "index", transform="growth"), "unknown transform"),
        (_units("real_gdp", "pct", transform="index"), "unknown transform"),
        (_units("real_gdp", "pct", **{"from": "garbage"}), "reads no measure"),
    ],
    ids=[
        "% change of debt",
        "% change of unemployment",
        "% change of a share",
        "GDP's % unchanged",
        "% as published of £",
        "£ of a share",
        "index as growth",
        "unknown transform",
        "no measure",
    ],
)
def test_a_unit_reads_only_a_measure_its_figures_are_in(variable, message):
    # The loader's rule (load.js, variableProblem): % as published reads a
    # measure in %, a % change one that is not, £ never one in %
    with pytest.raises(B.BuildError, match=message):
        B.check_units(variable)


def _with(variable_id: str, **fields):
    variable = next(v for v in V.VARIABLES if v.id == variable_id)
    return dataclasses.replace(variable, **fields)


def _level_with(variable_id: str, **fields):
    variable = next(v for v in V.VARIABLES if v.id == variable_id)
    return dataclasses.replace(
        variable, level=dataclasses.replace(variable.level, **fields)
    )


@pytest.mark.parametrize(
    "variable, message",
    [
        # Found by the twelfth review: a unit the page has no pill for, and
        # an index or % change of a share, built and then refused by the page
        (
            _with(
                "receipts",
                units={
                    **_units("receipts", "level").units,
                    "foo": {
                        "label": "Foo",
                        "axis": "£ billion",
                        "from": "level",
                    },
                },
            ),
            "unknown units",
        ),
        (_with("receipts", units={}), "unknown units"),
        (
            _units("receipts", "index", **{"from": "share"}),
            "taken of no level",
        ),
        (_units("receipts", "level", axis=""), "no label or axis"),
        (_units("receipts", "level", label=None), "no label or axis"),
        (_with("receipts", id="Receipts"), "no id, label or frequency"),
        (_with("receipts", label=""), "no id, label or frequency"),
        (_with("receipts", frequency="monthly"), "no id, label or frequency"),
        (_with("receipts", notes=("",)), "notes that are not text"),
        (_level_with("receipts", unit=""), "no unit or sources"),
        (_level_with("receipts", decimals=7), "no decimals"),
        (_level_with("receipts", decimals=-1), "no decimals"),
        (_level_with("receipts", decimals=1.5), "no decimals"),
        (_level_with("receipts", decimals=True), "no decimals"),
    ],
    ids=[
        "unknown unit",
        "no units",
        "index of a share",
        "no axis",
        "no label",
        "id unfit for an address",
        "no label",
        "unknown frequency",
        "blank note",
        "no unit",
        "seven decimals",
        "negative decimals",
        "fractional decimals",
        "decimals a flag",
    ],
)
def test_the_build_writes_no_variable_the_page_refuses(variable, message):
    # load.js's variableProblem and measureProblem, field by field
    with pytest.raises(B.BuildError, match=message):
        B.check_variable(variable)


def test_every_variable_is_one_the_page_takes():
    for variable in V.VARIABLES:
        B.check_variable(variable)


def test_every_committed_figure_is_written_to_four_decimals(committed):
    # The page refuses a figure that is not 0 and finer than 1e-4: one far
    # finer would underflow when a forecast is carried on (load.js, isFigure)
    for variable in committed["variables"]:
        for name in ("level", "share"):
            measure = variable.get(name)
            if not measure:
                continue
            for series in [measure["outturn"]] + [
                v["values"] for v in measure["vintages"]
            ]:
                for value in series.values():
                    assert value == round(value, B.DECIMALS)
                    assert value == 0 or abs(value) >= 10**-B.DECIMALS


@NO_DEADLINE
@given(st.dictionaries(st.integers(1900, 2100), st.floats(-1e9, 1e9)))
def test_the_build_writes_no_figure_finer_than_four_decimals(values):
    # Every series the file holds is written by _ordered (outturn and
    # forecasts alike): each figure to four decimals, the nearest to its
    # value, so 0 or at least 1e-4 in size
    written = B._ordered(values)
    assert list(written) == [str(p) for p in sorted(values)]
    for period, value in values.items():
        figure = written[str(period)]
        # The double nearest the value rounded to four decimals, worked in
        # exact decimals (half to even, as Python rounds)
        nearest = decimal.Decimal(value).quantize(
            decimal.Decimal(1).scaleb(-B.DECIMALS), decimal.ROUND_HALF_EVEN
        )
        assert figure == float(nearest)
        assert figure == 0 or abs(figure) >= 10**-B.DECIMALS


# --- An outturn check reads as true -------------------------------------------

CHECK_CASES = json.loads(
    (
        ROOT / "tests" / "fixtures" / "forecast_vintages_check_cases.json"
    ).read_text()
)["cases"]


def _apply_check_op(data: dict, op: dict) -> None:
    """One change from the shared check cases, made in place."""
    variable = next(v for v in data["variables"] if v["id"] == op["variable"])
    measure = variable[op["measure"]]
    if op["op"] == "set":
        measure["check"][op["field"]] = op["to"]
    elif op["op"] == "drop_check":
        del measure["check"]
    elif op["op"] == "kind":
        measure["kind"] = op["to"]
    else:
        raise ValueError(f"Unknown check case op {op['op']!r}")


def _check_check(data: dict, variable_id: str, name: str) -> None:
    variable = next(v for v in data["variables"] if v["id"] == variable_id)
    measure = variable[name]
    B.check_as_read(
        variable_id,
        variable["frequency"],
        measure["kind"],
        measure["unit"],
        measure.get("check"),
        measure["outturn"],
    )


@pytest.mark.parametrize("case", CHECK_CASES, ids=lambda c: c["name"])
def test_the_build_reads_checks_as_the_page_does(committed, case):
    # The page's loader (checkProblem) is held to the same cases in
    # src/utils/forecastVintages/check.test.js, so the two agree
    data = json.loads(json.dumps(committed))
    for op in case["ops"]:
        _apply_check_op(data, op)
    target = case["ops"][0]
    if case["accepted"]:
        _check_check(data, target["variable"], target["measure"])
    else:
        with pytest.raises(B.BuildError):
            _check_check(data, target["variable"], target["measure"])


def test_every_committed_check_reads_as_true(committed):
    for variable in committed["variables"]:
        for name in ("level", "share"):
            if variable.get(name):
                _check_check(committed, variable["id"], name)


# --- A series the page can draw -------------------------------------------------

SERIES_CASES = json.loads(
    (
        ROOT / "tests" / "fixtures" / "forecast_vintages_series_cases.json"
    ).read_text()
)["cases"]


@pytest.mark.parametrize("case", SERIES_CASES, ids=lambda c: c["name"])
def test_the_build_reads_series_as_the_page_does(case):
    # The page's loader (isSeries) is held to the same cases in
    # src/utils/forecastVintages/series.test.js, so the two agree
    values = json.loads(json.dumps(case["values"]))
    if case["accepted"]:
        B.series_as_read("a series", case["frequency"], values)
    else:
        with pytest.raises(B.BuildError):
            B.series_as_read("a series", case["frequency"], values)


def test_every_committed_series_reads_as_one(committed):
    for variable in committed["variables"]:
        for name in ("level", "share"):
            measure = variable.get(name)
            if not measure:
                continue
            for values in [measure["outturn"]] + [
                v["values"] for v in measure["vintages"]
            ]:
                B.series_as_read(variable["id"], variable["frequency"], values)


def test_an_outturn_with_a_year_missing_between_is_not_written(monkeypatch):
    # Found by the ninth review: the ONS reader leaves out a blank
    # observation, so a blank 1955 in ABMI left a gap the build wrote and
    # the page refused (the whole file, "real_gdp: no outturn")
    load_ons = B.load_ons

    def blank_1955(root, retrieved):
        series = load_ons(root, retrieved)
        del series["ABMI"].years[1955]
        return series

    monkeypatch.setattr(B, "load_ons", blank_1955)
    with pytest.raises(
        B.BuildError,
        match="real_gdp: its outturn has no value between 1954 and 1956",
    ):
        B.build(ROOT)


# --- A file the page takes as a whole -----------------------------------------

DOCUMENT_CASES = json.loads(
    (
        ROOT / "tests" / "fixtures" / "forecast_vintages_document_cases.json"
    ).read_text()
)["cases"]


def _apply_document_op(data: dict, op: dict) -> None:
    """One change from the shared document cases, made in place."""
    parent = data
    for step in op["path"][:-1]:
        parent = parent[step]
    last = op["path"][-1]
    if op["op"] == "set":
        parent[last] = json.loads(json.dumps(op["to"]))
    elif op["op"] == "delete":
        del parent[last]
    elif op["op"] == "append":
        parent[last].append(json.loads(json.dumps(op["to"])))
    else:
        raise ValueError(f"Unknown document case op {op['op']!r}")


@pytest.mark.parametrize("case", DOCUMENT_CASES, ids=lambda c: c["name"])
def test_the_build_reads_the_file_as_the_page_does(committed, case):
    # The page's loader (forecastVintagesProblem) is held to the same cases
    # in src/utils/forecastVintages/document.test.js, so the two agree
    data = json.loads(json.dumps(committed))
    for op in case["ops"]:
        _apply_document_op(data, op)
    if case["accepted"]:
        B.check_document(data)
    else:
        with pytest.raises(B.BuildError):
            B.check_document(data)


def test_the_build_writes_no_variable_twice(monkeypatch):
    # Found by the thirteenth review: each variable's id was checked on its
    # own, so the build wrote receipts as a second "spending" and the page
    # refused the whole file ("a variable twice")
    variables = [
        dataclasses.replace(v, id="spending") if v.id == "receipts" else v
        for v in V.VARIABLES
    ]
    monkeypatch.setattr(V, "VARIABLES", variables)
    with pytest.raises(B.BuildError, match="a variable twice"):
        B.build(ROOT)


def test_the_build_checks_the_file_as_json_reads_it(monkeypatch):
    # The page reads the JSON the build writes, not the objects it holds
    seen = []
    monkeypatch.setattr(B, "check_document", seen.append)
    built = B.build(ROOT)
    assert len(seen) == 1 and seen[0] is not built
    assert seen[0] == json.loads(B.output_text(built))


def test_each_measure_names_exactly_the_files_it_is_built_from(built):
    # Found by the fourteenth review: Sources could leave out a file a
    # measure is built from, under a card saying each file is listed. Each
    # measure names its files as the build reads them: from those alone
    # its outturn is the one written, and each one named is read
    retrieved = B.read_retrieved(ROOT)
    series, hpi = B.load_ons(ROOT, retrieved), B.load_hpi(ROOT, retrieved)
    efo_ids = [file.id for file in sources.EFO_TABLES]
    named = 0
    for variable in V.VARIABLES:
        record = next(v for v in built["variables"] if v["id"] == variable.id)
        for name in ("level", "share"):
            measure = getattr(variable, name)
            if measure is None:
                continue
            files = record[name]["files"]
            only = {cdid: series[cdid] for cdid in files["ons"]}
            whole = B._ordered(measure.outturn(series, hpi))
            assert (
                B._ordered(measure.outturn(only, hpi if files["hpi"] else {}))
                == whole
            )
            for cdid in files["ons"]:
                without = {k: v for k, v in only.items() if k != cdid}
                with pytest.raises(KeyError):
                    measure.outturn(without, hpi)
            if files["hpi"]:
                assert measure.outturn(only, {}) == {}
            assert files["obr"] == (
                efo_ids if measure.sheet is None else [sources.HOFD.id]
            )
            named += len(files["ons"]) + files["hpi"]
    assert named > 20
    by_id = {v["id"]: v for v in built["variables"]}
    assert by_id["receipts"]["share"]["files"]["ons"] == [
        "ANBT",
        "BKTL",
        "IL65",
    ]
    assert by_id["house_prices"]["level"]["files"] == {
        "obr": [sources.HOFD.id],
        "ons": [],
        "hpi": True,
    }


def test_the_economy_tables_hold_each_forecast_once(monkeypatch):
    # Found by the fourteenth review: a second March 2026 entry (a
    # re-issue, say) passed the date-order check and was written twice
    reissue = dataclasses.replace(
        sources.EFO_TABLES[-1], id="efo_march_2026_reissue"
    )
    monkeypatch.setattr(sources, "EFO_TABLES", sources.EFO_TABLES + (reissue,))
    with pytest.raises(B.BuildError, match="each forecast once"):
        B.read_efo_tables(ROOT)


def test_the_build_writes_no_forecast_twice(monkeypatch):
    # However the tables come to hold one twice, the build checks every
    # measure's forecasts as the page does before writing them
    read = B.read_efo_tables
    monkeypatch.setattr(
        B, "read_efo_tables", lambda root: (tables := read(root)) + tables[-1:]
    )
    with pytest.raises(
        B.BuildError,
        match="rhdi_per_person .level.: the March 2026 forecast twice",
    ):
        B.build(ROOT)


# --- A forecast's figures after its launch -------------------------------------

COMMITTED_TEXT = COMMITTED.read_text()


def _in_bounds(kind: str):
    """Figures a forecast of `kind` may hold after its launch, any of them."""
    if kind == "growth":
        return st.floats(-99.999999, 99.999999)
    if kind == "index":
        return st.floats(1e-6, 1e6)
    return st.floats(-1e9, 1e9)


def _kind_verdict(data: dict, variable_id: str, name: str) -> bool:
    try:
        _kind_check(data, variable_id, name)
    except B.BuildError:
        return False
    return True


@NO_DEADLINE
@given(st.sampled_from(KIND_CASES), st.integers(0, 64), st.data())
def test_figures_after_the_launch_never_change_the_verdict(case, which, draw):
    # From any shared case, taken or refused: two sets of figures after a
    # forecast's launch, each within the bounds every figure keeps (a
    # growth rate under 100% either way, an index above zero), give the
    # same verdict
    target = case["ops"][0]
    verdicts = set()
    for _ in range(2):
        data = json.loads(COMMITTED_TEXT)
        for op in case["ops"]:
            _apply_kind_op(data, op)
        variable = next(
            v for v in data["variables"] if v["id"] == target["variable"]
        )
        measure = variable[target["measure"]]
        vintage = measure["vintages"][which % len(measure["vintages"])]
        for key in vintage["values"]:
            if key != str(vintage["launch"]):
                vintage["values"][key] = draw.draw(_in_bounds(measure["kind"]))
        verdicts.add(
            _kind_verdict(data, target["variable"], target["measure"])
        )
    assert len(verdicts) == 1


@NO_DEADLINE
@given(
    st.sampled_from(
        [m for m in _committed_measures() if m[0] in ("cpi", "real_gdp")]
        + [("rhdi_per_person", "level")]
    ),
    st.integers(0, 64),
    st.data(),
)
def test_a_figure_past_its_bounds_is_always_refused(target, which, draw):
    # A growth rate of 100% or more either way, or an index at or below
    # zero, wherever it falls in a forecast
    data = json.loads(COMMITTED_TEXT)
    variable_id, name = target
    variable = next(v for v in data["variables"] if v["id"] == variable_id)
    measure = variable[name]
    vintage = measure["vintages"][which % len(measure["vintages"])]
    key = draw.draw(st.sampled_from(sorted(vintage["values"])))
    if measure["kind"] == "growth":
        past = st.floats(100, 1e9) | st.floats(-1e9, -100)
    else:
        past = st.floats(-1e9, 0)
    vintage["values"][key] = draw.draw(past)
    assert not _kind_verdict(data, variable_id, name)


@NO_DEADLINE
@given(
    st.sampled_from(
        [
            (v["id"], name)
            for v in json.loads(COMMITTED_TEXT)["variables"]
            for name in ("level", "share")
            if v.get(name) and v[name]["kind"] != "index"
        ]
    ),
    st.lists(st.booleans(), min_size=40, max_size=40),
)
def test_forecasts_may_start_the_period_after_their_launch(target, drop):
    # As the build lets a level or growth forecast: whichever drop their
    # launch figure, while the first keeps it, the file is still written
    data = json.loads(COMMITTED_TEXT)
    variable_id, name = target
    variable = next(v for v in data["variables"] if v["id"] == variable_id)
    for i, vintage in enumerate(variable[name]["vintages"]):
        if i > 0 and drop[i % len(drop)]:
            vintage["values"].pop(str(vintage["launch"]), None)
    _kind_check(data, variable_id, name)


# --- Downloads never reach obr.uk --------------------------------------------


def test_a_redirect_to_obr_uk_is_never_followed(tmp_path):
    seen = []

    def handler(request):
        seen.append(request.url.host)
        # An archive capture that redirects to the live page
        return httpx.Response(
            302, headers={"Location": "https://obr.uk/docs/dlm_uploads/x.xlsx"}
        )

    with fetch.make_client(httpx.MockTransport(handler)) as http:
        with pytest.raises(fetch.FetchError, match="obr.uk is never fetched"):
            fetch.fetch_archived(http, tmp_path, log=lambda _: None)
    assert seen == ["web.archive.org"]


def test_obr_uk_is_never_requested_directly():
    seen = []

    def handler(request):
        seen.append(request.url.host)
        return httpx.Response(200, content=b"")

    with fetch.make_client(httpx.MockTransport(handler)) as http:
        for url in ("https://obr.uk/x", "https://www.obr.uk/x"):
            with pytest.raises(fetch.FetchError):
                fetch._get(http, url)
    assert seen == []


# obr.uk however DNS would read it: any case, a final dot (the review found
# "obr.uk." passed the first hook), and any subdomain
_OBR_SPELLINGS = (
    "https://obr.uk./x",
    "https://www.obr.uk./x",
    "https://OBR.UK/x",
    "https://Www.Obr.Uk./docs/x.xlsx",
    "http://obr.uk:443/x",
)


@pytest.mark.parametrize("url", _OBR_SPELLINGS)
def test_obr_uk_is_refused_however_its_host_is_written(url):
    seen = []

    def handler(request):
        seen.append(request.url.host)
        return httpx.Response(200, content=b"")

    with fetch.make_client(httpx.MockTransport(handler)) as http:
        with pytest.raises(fetch.FetchError, match="obr.uk is never fetched"):
            fetch._get(http, url)
    assert seen == []


@pytest.mark.parametrize("location", _OBR_SPELLINGS)
def test_a_redirect_to_obr_uk_however_written_is_never_followed(location):
    seen = []

    def handler(request):
        seen.append(request.url.host)
        return httpx.Response(302, headers={"Location": location})

    with fetch.make_client(httpx.MockTransport(handler)) as http:
        with pytest.raises(fetch.FetchError, match="obr.uk is never fetched"):
            fetch._get(http, "https://web.archive.org/web/2025id_/x")
    assert seen == ["web.archive.org"]


_LABEL = st.text(
    "abcdefghijklmnopqrstuvwxyz0123456789-", min_size=1, max_size=8
)


@NO_DEADLINE
@given(
    st.lists(_LABEL, max_size=3),
    st.lists(st.booleans(), min_size=6, max_size=6),
    st.booleans(),
)
def test_no_spelling_of_an_obr_uk_host_is_ever_requested(labels, upper, dot):
    name = ".".join([*labels, "obr", "uk"])
    host = "".join(
        c.upper() if upper[i % len(upper)] else c for i, c in enumerate(name)
    )
    seen = []

    def handler(request):
        seen.append(request.url.host)
        return httpx.Response(200, content=b"")

    with fetch.make_client(httpx.MockTransport(handler)) as http:
        with pytest.raises(fetch.FetchError, match="obr.uk is never fetched"):
            fetch._get(http, f"https://{host}{'.' if dot else ''}/x")
    assert seen == []


@NO_DEADLINE
@given(st.lists(_LABEL, min_size=1, max_size=4))
def test_only_the_hosts_the_inputs_come_from_are_requested(labels):
    host = ".".join(labels)
    seen = []

    def handler(request):
        seen.append(request.url.host)
        return httpx.Response(200, content=b"ok")

    with fetch.make_client(httpx.MockTransport(handler)) as http:
        if host.rstrip(".") in fetch.SOURCE_HOSTS:
            assert fetch._get(http, f"https://{host}/x") == b"ok"
        else:
            with pytest.raises(fetch.FetchError):
                fetch._get(http, f"https://{host}/x")
            assert seen == []


# Found by review round 22: hosts httpx cannot read raised its InvalidURL,
# which is not an httpx.HTTPError, and the property above drew them about
# one run in twenty
@pytest.mark.parametrize(
    "url",
    ["https://0.0.0.00/x", "https://0.0.0.0000/x", "https://1.2.3.04/x"],
)
def test_a_url_httpx_cannot_read_is_refused_unrequested(url):
    seen = []

    def handler(request):
        seen.append(request.url.host)
        return httpx.Response(200, content=b"ok")

    with fetch.make_client(httpx.MockTransport(handler)) as http:
        with pytest.raises(fetch.FetchError, match="Invalid IPv4 address"):
            fetch._get(http, url)
    assert seen == []


def test_a_redirect_off_the_source_hosts_is_refused():
    seen = []

    def handler(request):
        seen.append(request.url.host)
        return httpx.Response(
            302, headers={"Location": "https://example.com/data.json"}
        )

    with fetch.make_client(httpx.MockTransport(handler)) as http:
        with pytest.raises(fetch.FetchError, match="not a host the inputs"):
            fetch._get(http, sources.ONS["ABMI"].url)
    assert seen == ["www.ons.gov.uk"]


def test_every_download_is_from_a_source_host():
    # fetch_all asks for these, so none may be refused
    urls = [s.url for s in sources.ONS.values()]
    urls.append(sources.HPI_URL.format(to="2026-01-01"))
    urls += [
        f.archive_url
        for f in (sources.HOFD,) + sources.EFO_TABLES
        if f.archive_url
    ]
    assert {httpx.URL(u).host for u in urls} == fetch.SOURCE_HOSTS


# --- Periods and launches ----------------------------------------------------


@NO_DEADLINE
@given(st.integers(1900, 2100), st.integers(1, 4))
def test_quarter_numbers_round_trip(year, quarter):
    label = periods.quarter_label(year, quarter)
    assert periods.quarter_from_number(periods.quarter_number(label)) == label
    following = periods.next_period(label)
    assert periods.sort_key(following) == periods.sort_key(label) + 1


@NO_DEADLINE
@given(st.integers(1900, 2100))
def test_years_and_fiscal_years_step_by_one(year):
    assert periods.next_period(year) == year + 1
    fy = periods.fiscal_label(year)
    assert periods.next_period(fy) == periods.fiscal_label(year + 1)
    assert (
        periods.sort_key(periods.next_period(fy)) == periods.sort_key(fy) + 1
    )


@NO_DEADLINE
@given(st.integers(2010, 2030), st.integers(1, 12))
def test_every_launch_is_complete_before_publication(year, month):
    published_quarter = year * 4 + (month - 1) // 3
    q = periods.quarter_number(periods.quarterly_launch(year, month))
    assert q == published_quarter - 2
    assert periods.calendar_launch(year, month) == year - 1
    # A fiscal year ends in March of its second year: the launch year has
    # ended before the month of publication, and the one after it has not.
    start = int(periods.fiscal_launch(year, month)[:4])
    assert (start + 1, 3) < (year, month)
    assert not (start + 2, 3) < (year, month)


def test_launches_of_known_forecasts():
    assert periods.launch("quarterly", 2026, 3) == "2025Q3"
    assert periods.launch("quarterly", 2025, 11) == "2025Q2"
    assert periods.launch("calendar", 2026, 3) == 2025
    assert periods.launch("fiscal", 2026, 3) == "2024-25"
    assert periods.launch("fiscal", 2025, 11) == "2024-25"
    assert periods.launch("fiscal", 2010, 6) == "2009-10"
    with pytest.raises(ValueError):
        periods.launch("monthly", 2026, 3)


# --- Command line -----------------------------------------------------------


def test_the_cli_parses_the_command():
    parsed = parse_args(["forecast-vintages", "--reviewer", ""])
    assert parsed.command == "forecast-vintages"
    assert parsed.fetch is False
    assert parsed.output == B.OUTPUT_PATH
    assert parse_args(["forecast-vintages", "--fetch"]).fetch is True
